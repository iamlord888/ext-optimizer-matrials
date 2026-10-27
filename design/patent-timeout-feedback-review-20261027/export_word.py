"""Export the current patent manuscript as an editable Word document.

Run with the Codex bundled Python. The paired Markdown is the content source;
this helper keeps the requested Chinese typography consistent on regeneration.
"""

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree


FONT_NAME = "宋体"
BODY_SIZE = 12
REPO = Path(__file__).resolve().parents[2]
PUBLIC_BASE = "https://github.com/iamlord888/ext-optimizer-matrials/blob/main/"
INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*\n]+\*|\[[^\]]+\]\([^)]+\))")
# Only the manuscript's mathematical symbols use subscript notation. Ordinary
# code identifiers such as plan_node_id and get_next retain their underscores.
MATH_SUBSCRIPT = re.compile(
    r"(?<![A-Za-z0-9_])([ABCEPTt])"
    r"(?:_\{(i(?:[+-][0-9]+)?|[0-9]+|best|cap)\}"
    r"|_(i|[0-9]+|best|cap)|([₀₁₂₃₄₅₆₇₈₉ᵢ]))"
    r"(?![A-Za-z0-9_])"
)
SUBSCRIPT_CHARACTERS = str.maketrans("₀₁₂₃₄₅₆₇₈₉ᵢ", "0123456789i")


def apply_font(properties, size=BODY_SIZE, bold=False, italic=False):
    """Set every Word script slot explicitly, avoiding theme-font overrides."""
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), FONT_NAME)
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + key), None)
    for name, value in (("sz", str(size * 2)), ("szCs", str(size * 2)), ("color", "000000")):
        item = properties.find(qn("w:" + name))
        if item is None:
            item = OxmlElement("w:" + name)
            properties.append(item)
        item.set(qn("w:val"), value)
    for name, enabled in (("b", bold), ("i", italic)):
        item = OxmlElement("w:" + name)
        item.set(qn("w:val"), "1" if enabled else "0")
        properties.append(item)


def format_paragraph(paragraph, indent=True):
    """Use true 1.5-line spacing and a two-character first-line indent."""
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(2 * BODY_SIZE if indent else 0)
    fmt.widow_control = True
    props = paragraph._p.get_or_add_pPr()
    ind = props.find(qn("w:ind"))
    ind.set(qn("w:firstLineChars"), "200" if indent else "0")
    for name in ("autoSpaceDE", "autoSpaceDN"):
        item = OxmlElement("w:" + name)
        item.set(qn("w:val"), "0")
        props.append(item)


def link_target(value, source):
    """Make local source references usable outside this Mac workspace."""
    value = value.strip("<>")
    if value.startswith(("https://", "http://", "mailto:")):
        return value
    target = (Path(value) if value.startswith("/") else source.parent / value).resolve()
    if target.is_relative_to(REPO):
        relative = target.relative_to(REPO).as_posix()
        if relative.startswith(".scratch/patience5-ab-20260928/"):
            relative = "data/" + relative[len(".scratch/"):]
        return PUBLIC_BASE + relative
    if target.name in ("P019.pdf", "1601.05748v1.pdf"):
        return PUBLIC_BASE + "sources/papers/" + target.name
    return target.as_uri()


def add_run(paragraph, value, bold=False, italic=False):
    """Use native text subscripts so indices flow with Chinese punctuation."""
    position = 0
    for match in MATH_SUBSCRIPT.finditer(value):
        suffix = next(item for item in match.groups()[1:] if item is not None)
        for text, is_subscript in ((value[position:match.start()], False), (match[1], False), (suffix.translate(SUBSCRIPT_CHARACTERS), True)):
            if text:
                run = paragraph.add_run(text)
                apply_font(run._r.get_or_add_rPr(), bold=bold, italic=italic)
                if is_subscript:
                    run.font.subscript = True
        position = match.end()
    if position < len(value):
        run = paragraph.add_run(value[position:])
        apply_font(run._r.get_or_add_rPr(), bold=bold, italic=italic)


def add_inline(paragraph, value, source):
    """Preserve wording while translating Markdown emphasis and links."""
    position = 0
    for match in INLINE.finditer(value):
        add_run(paragraph, value[position:match.start()])
        token = match.group()
        if token.startswith("`"):
            add_run(paragraph, token[1:-1])
        elif token.startswith("**"):
            add_run(paragraph, token[2:-2], bold=True)
        elif token.startswith("*"):
            add_run(paragraph, token[1:-1], italic=True)
        else:
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            label = link[1].replace("`", "")
            relation = paragraph.part.relate_to(link_target(link[2], source), RT.HYPERLINK, is_external=True)
            element = OxmlElement("w:hyperlink")
            element.set(qn("r:id"), relation)
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            apply_font(properties)
            run.append(properties)
            text = OxmlElement("w:t")
            text.text = label
            run.append(text)
            element.append(run)
            paragraph._p.append(element)
        position = match.end()
    add_run(paragraph, value[position:])


def math_run(value):
    run = OxmlElement("m:r")
    props = OxmlElement("m:rPr")
    plain = OxmlElement("m:sty")
    plain.set(qn("m:val"), "p")
    props.append(plain)
    run.append(props)
    word_props = OxmlElement("w:rPr")
    apply_font(word_props)
    run.append(word_props)
    text = OxmlElement("m:t")
    text.text = value
    run.append(text)
    return run


def subscript(base, suffix):
    item = OxmlElement("m:sSub")
    expression = OxmlElement("m:e")
    expression.append(math_run(base))
    sub = OxmlElement("m:sub")
    sub.append(math_run(suffix))
    item.extend((expression, sub))
    return item


def add_timeout_equation(doc, value):
    """Create editable Word math for the manuscript's single display formula."""
    compact = re.sub(r"\s+", "", value)
    if compact != "T_cap=min(T₀,1.25×t_best)":
        raise ValueError("Unsupported display equation; preserve it explicitly before exporting: " + value)
    paragraph = doc.add_paragraph()
    format_paragraph(paragraph, indent=False)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    container = OxmlElement("m:oMathPara")
    equation = OxmlElement("m:oMath")
    equation.extend((subscript("T", "cap"), math_run(" = min("), subscript("T", "0"), math_run(", 1.25 × "), subscript("t", "best"), math_run(")")))
    container.append(equation)
    paragraph._p.append(container)


def create_document(source, output):
    content = source.read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.9)
    section.left_margin = section.right_margin = Inches(1)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = FONT_NAME
        style.font.size = Pt(BODY_SIZE)
        style.font.color.rgb = RGBColor(0, 0, 0)
        apply_font(style.element.get_or_add_rPr())
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(0)
    # The bundled Word starter carries a blue Title border; a patent
    # disclosure uses only typography and whitespace for its heading ladder.
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    font_part = doc.part.part_related_by(RT.FONT_TABLE)
    fonts = etree.fromstring(font_part.blob)
    entry = OxmlElement("w:font")
    entry.set(qn("w:name"), FONT_NAME)
    for name, value in (("altName", "Songti SC"), ("family", "roman"), ("charset", "86")):
        item = OxmlElement("w:" + name)
        item.set(qn("w:val"), value)
        entry.append(item)
    fonts.append(entry)
    font_part._blob = etree.tostring(fonts, xml_declaration=True, encoding="UTF-8", standalone=True)
    lines = content.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("```"):
            block = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            add_timeout_equation(doc, "\n".join(block))
            index += 1
            continue
        heading = re.match(r"^(#{1,4}) (.+)", line)
        if heading:
            level, title = len(heading[1]), heading[2]
            paragraph = doc.add_paragraph(style="Title" if level == 1 else "Heading " + str(level - 1))
            format_paragraph(paragraph, indent=False)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(12 if level > 1 else 0)
            paragraph.paragraph_format.space_after = Pt(6 if level > 1 else 12)
            run = paragraph.add_run(title)
            apply_font(run._r.get_or_add_rPr(), size={1: 16, 2: 14, 3: 13, 4: 12}[level], bold=True)
            if level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.core_properties.title = title
            index += 1
            continue
        if line.startswith(("|", "![")):
            raise ValueError("This manuscript exporter requires explicit handling of added tables or figures.")
        paragraph = doc.add_paragraph()
        format_paragraph(paragraph)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(paragraph, line[2:] if line.startswith("- ") else line, source)
        index += 1
    footer = section.footer.paragraphs[0]
    format_paragraph(footer, indent=False)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    apply_font(props, size=10)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    footer._p.append(field)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    inline_subscripts = sum(1 for item in doc.element.iter(qn("w:vertAlign")) if item.get(qn("w:val")) == "subscript")
    manifest = {"source": str(source), "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "output": str(output), "output_sha256": hashlib.sha256(output.read_bytes()).hexdigest(), "font": FONT_NAME, "body_size_pt": BODY_SIZE, "line_spacing": 1.5, "body_first_line_characters": 2, "paragraphs": len(doc.paragraphs), "editable_display_equations": 1, "inline_subscript_runs": inline_subscripts, "native_subscripts": inline_subscripts + sum(1 for item in doc.element.iter(qn("m:sSub")))}
    output.with_suffix(".word-sync.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=Path(__file__).with_name("revised-manuscript.md"))
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    source = args.source.resolve()
    create_document(source, (args.output or source.with_suffix(".docx")).resolve())
