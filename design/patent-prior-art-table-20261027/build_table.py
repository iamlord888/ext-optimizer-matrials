"""Build the literature register from reviewed bibliographic records.

Reuse the existing disclosure exporter for typography and paragraph handling.
Use an A4 landscape page for the four requested fields at 12pt, rather than
shrinking the text to fit a portrait page.
"""

import contextlib
import hashlib
import html
import importlib.util
import io
import json
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt


ROOT = Path(__file__).resolve().parent
HELPER = ROOT.parent / "patent-timeout-feedback-review-20261027" / "export_word.py"
SPEC = importlib.util.spec_from_file_location("patent_word_export", HELPER)
WORD = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(WORD)


def display_date(value):
    """Keep each date and its publication-stage label on readable lines."""
    if value == "首次公开日期未载明":
        return "首次公开日期\n未载明"
    return value.replace("；", "\n").replace("（", "\n").replace("）", "")


def record_cells(record):
    """Keep document identity, dates, source links and exact locations together."""
    identity = f"**{record['english_title']}**\n{record['publication']}"
    sources = "\n".join(f"[{item['url']}]({item['url']})" for item in record["sources"][:1] if item["url"])
    return [identity, display_date(record["date"]), sources, record["locations"]]


def build_word(data):
    """Insert one editable, repeating-header table into the shared Word layout."""
    with tempfile.TemporaryDirectory(prefix="patent-literature-") as temporary:
        temporary = Path(temporary)
        source = temporary / "register.md"
        source.write_text("# " + data["title"] + "\n\n" + data["scope"] + "\n", encoding="utf-8")
        draft = temporary / "register.docx"
        with contextlib.redirect_stdout(io.StringIO()):
            WORD.create_document(source, draft)
        document = Document(draft)
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Mm(297), Mm(210)
        section.top_margin = section.bottom_margin = Inches(0.55)
        section.left_margin = section.right_margin = Inches(0.55)
        anchor = document.paragraphs[-1]
        cells = [data["headers"], *[record_cells(record) for record in data["records"]]]
        table = document.add_table(rows=len(cells), cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.style = "Table Grid"
        widths = (4.5, 1.2, 2.5, 2.3)
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for index, (row, values) in enumerate(zip(table.rows, cells)):
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement("w:cantSplit"))
            if index == 0:
                properties.append(OxmlElement("w:tblHeader"))
            for cell, width, value in zip(row.cells, widths, values):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                paragraph = cell.paragraphs[0]
                WORD.format_paragraph(paragraph, indent=False)
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.space_after = Pt(3)
                if index == 0:
                    run = paragraph.add_run(value)
                    WORD.apply_font(run._r.get_or_add_rPr(), bold=True)
                    fill = OxmlElement("w:shd")
                    fill.set(qn("w:fill"), "F2F2F2")
                    cell._tc.get_or_add_tcPr().append(fill)
                else:
                    WORD.add_inline(paragraph, value, ROOT / "literature.md")
        anchor._p.addnext(table._tbl)
        output = ROOT / "literature.docx"
        document.save(output)
        return output


def build_markdown(data):
    """Export a readable source table with the same records as the Word file."""
    lines = ["# " + data["title"], "", data["scope"], "", "| " + " | ".join(data["headers"]) + " |", "| " + " | ".join(["---"] * 4) + " |"]
    for record in data["records"]:
        lines.append("| " + " | ".join(cell.replace("\n", "<br>") for cell in record_cells(record)) + " |")
    lines.append("")
    output = ROOT / "literature.md"
    output.write_text("\n".join(lines), encoding="utf-8")
    return output


def build_html(data):
    """Write literal table markup with linked primary sources."""
    rows = []
    esc = html.escape
    for record in data["records"]:
        identity = f'<strong>{esc(record["english_title"])}</strong><p>{esc(record["publication"])}</p>'
        dates = esc(display_date(record["date"])).replace("\n", "<br>")
        sources = '<br>'.join(f'<a href="{esc(item["url"], quote=True)}" target="_blank" rel="noreferrer">{esc(item["url"])}</a>' for item in record['sources'][:1] if item['url'])
        rows.append('<tr>' + ''.join('<td>' + value + '</td>' for value in [identity, dates, sources, esc(record['locations'])]) + '</tr>')
    headers = '<tr>' + ''.join('<th scope="col">' + esc(value) + '</th>' for value in data['headers']) + '</tr>'
    page = '''<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><link rel="icon" href="data:,"><title>现有技术文献信息表</title><style>
doc-page:not(:defined){visibility:hidden}body{margin:0}doc-page{font-family:"Songti SC","SimSun",serif;font-size:16px;line-height:1.75;color:#202020}h1{font-size:27px;line-height:1.4;margin:0 0 18px}h2{font-size:21px;margin:32px 0 14px}p{margin:12px 0;overflow-wrap:break-word}nav{display:flex;flex-wrap:wrap;gap:12px 24px;margin-bottom:24px;font-family:-apple-system,"PingFang SC",sans-serif;font-size:14px}a{color:#315b78;text-underline-offset:4px}table{width:100%;border-collapse:collapse;table-layout:fixed;font-size:15px;line-height:1.65}th,td{padding:11px;border:1px solid #c8c8c8;text-align:left;vertical-align:top;overflow-wrap:anywhere}th{background:#f2f2f2;font-weight:700}td p{margin:8px 0 0}th:nth-child(1){width:37%}th:nth-child(2){width:14%}th:nth-child(3){width:26%}th:nth-child(4){width:23%}@media screen and (max-width:1250px){doc-page{min-width:0;padding:24px 16px;--doc-page-w:min(12in,calc(100vw - 32px))!important;--doc-page-margin:24px!important}}@media screen and (max-width:600px){doc-page{font-size:15px}table{font-size:13px}th,td{padding:7px}h1{font-size:23px}}@media print{nav{display:none}}
</style></head><body><doc-page orientation="landscape" margin="0.55in"><h1>''' + esc(data['title']) + '''</h1><nav aria-label="文档操作"><a href="literature.docx" download>下载Word</a><a href="literature.md" download>下载文字表格</a><a href="../patent-timeout-feedback-terminology-20261027/Terms.html">缩略语与关键术语</a></nav><p>''' + esc(data['scope']) + '</p><table><thead>' + headers + '</thead><tbody>' + '\n'.join(rows) + '</tbody></table></doc-page><script src="doc-page.js"></script></body></html>\n'
    output = ROOT / "Literature.html"
    output.write_text(page, encoding="utf-8")
    return output


if __name__ == "__main__":
    data = json.loads((ROOT / "literature.json").read_text(encoding="utf-8"))
    if len(data["records"]) != 13:
        raise ValueError("Expected the thirteen distinct documents reviewed for this manuscript")
    word = build_word(data)
    markdown = build_markdown(data)
    preview = build_html(data)
    print(json.dumps({"records": len(data['records']), "word": str(word), "markdown": str(markdown), "html": str(preview), "word_sha256": hashlib.sha256(word.read_bytes()).hexdigest()}, ensure_ascii=False, indent=2))
