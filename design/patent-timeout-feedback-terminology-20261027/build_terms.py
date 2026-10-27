"""Build the paired terminology Word and HTML from terms.md.

Reuse the manuscript's font, paragraph, inline text and page setup helpers.
The table is authored as a native editable Word table.
"""

import contextlib
import hashlib
import html
import importlib.util
import io
import json
import re
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "terms.md"
HELPER = ROOT.parent / "patent-timeout-feedback-review-20261027" / "export_word.py"
SPEC = importlib.util.spec_from_file_location("patent_word_export", HELPER)
SHARED = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(SHARED)


def parse_source():
    """Read the explicitly scoped table and ordinary paragraph manuscript."""
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    table = []
    for line in lines:
        if line.startswith("|"):
            cells = [part.strip() for part in line.strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                table.append(cells)
    if len(table) != 8 or any(len(row) != 3 for row in table):
        raise ValueError("Expected the reviewed three-column, seven-entry abbreviation table")
    return lines, table


def build_word(lines, rows):
    """Reuse the existing disclosure exporter and insert the reviewed table."""
    with tempfile.TemporaryDirectory(prefix="patent-terminology-") as temporary:
        temp_root = Path(temporary)
        text_source = temp_root / "terms-text.md"
        text_source.write_text("\n".join(line for line in lines if not line.startswith("|")) + "\n", encoding="utf-8")
        draft = temp_root / "terms-text.docx"
        with contextlib.redirect_stdout(io.StringIO()):
            SHARED.create_document(text_source, draft)
        doc = Document(draft)
        table_anchor = next(paragraph for paragraph in doc.paragraphs if paragraph.text == "缩略语表")
        terms_heading = next(paragraph for paragraph in doc.paragraphs if paragraph.text == "关键术语说明")
        terms_heading.paragraph_format.page_break_before = True
        table = doc.add_table(rows=len(rows), cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.style = "Table Grid"
        widths = (2.0, 3.35, 1.15)
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for row_index, (row, values) in enumerate(zip(table.rows, rows)):
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                properties.append(OxmlElement("w:tblHeader"))
            for cell, width, value in zip(row.cells, widths, values):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                SHARED.format_paragraph(paragraph, indent=False)
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.space_after = Pt(3)
                run = paragraph.add_run(value)
                SHARED.apply_font(run._r.get_or_add_rPr(), bold=row_index == 0)
                if row_index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "F2F2F2")
                    cell._tc.get_or_add_tcPr().append(shading)
        table_anchor._p.addnext(table._tbl)
        output = ROOT / "terms.docx"
        doc.save(output)
        return output, len(doc.paragraphs)


def inline_html(value):
    """Preserve literal source text and the manuscript's simple emphasis."""
    value = html.escape(value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", value)
    return re.sub(r"`([^`]+)`", r"<code>\1</code>", value)


def build_html(lines, rows):
    content = []
    table_written = False
    for line in lines:
        if not line.strip():
            continue
        if line.startswith("|"):
            if not table_written:
                content.append('<table><thead><tr>' + ''.join('<th scope="col">' + html.escape(cell) + '</th>' for cell in rows[0]) + '</tr></thead><tbody>')
                content.extend('<tr>' + ''.join('<td>' + html.escape(cell) + '</td>' for cell in row) + '</tr>' for row in rows[1:])
                content.append('</tbody></table>')
                table_written = True
            continue
        heading = re.match(r"^(#{1,2}) (.+)$", line)
        if heading:
            level = len(heading[1])
            name = heading[2]
            section_id = 'document-title' if level == 1 else ('terms' if name == '关键术语说明' else 'abbreviations')
            content.append(f'<h{level} id="{section_id}">{html.escape(name)}</h{level}>')
            if level == 1:
                content.append('<nav aria-label="文档操作"><a href="terms.docx" download>下载Word</a><a href="terms.md" download>下载文字稿</a><a href="#terms">关键术语</a></nav>')
        else:
            content.append('<p>' + inline_html(line) + '</p>')
    notes = '<details><summary>收录范围说明</summary><p>缩略语表按当前整理稿正文实际用词收录。Rows Hint和Outline属于技术用语，EXPLAIN与EXPLAIN ANALYZE属于数据库命令，均在关键术语中解释。</p><p>QueryDesc、Plan、PlanState、Instrumentation及get_next继续在实施例二随文说明。active与terminal的含义归入本地EOF；EXACT、LOWER_BOUND和NONE是反馈分类标签。本文对部分反馈、反馈目标等概念所给英文为说明性译名，不另造缩写。</p><p>基数估算值与调优工具保存的输入值应在实施例中分别定义；图中的A-rows、E-rows及数学变量宜在图注或符号说明中解释。</p><p><a href="https://gemini.google.com/app/a7d6021460243a25" target="_blank" rel="noreferrer">查看与Gemini的讨论</a></p></details>'
    page = '''<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><link rel="icon" href="data:,"><title>专利交底书缩略语和关键术语</title>
<style>
doc-page:not(:defined){visibility:hidden}doc-page{font-family:"Songti SC","SimSun",serif;color:#202020;font-size:16px;line-height:1.8}h1{font-size:27px;line-height:1.4;font-weight:700;margin:0 0 18px}h2{font-size:21px;line-height:1.5;margin:38px 0 14px}p{margin:14px 0;text-align:justify;overflow-wrap:break-word}nav{display:flex;flex-wrap:wrap;gap:12px 24px;font-family:-apple-system,"PingFang SC",sans-serif;font-size:14px;margin-bottom:30px}a{color:#315b78;text-underline-offset:4px}table{width:100%;border-collapse:collapse;table-layout:fixed;font-size:15px;line-height:1.65}th,td{padding:10px 11px;border:1px solid #c8c8c8;text-align:left;vertical-align:top;overflow-wrap:break-word}th{font-weight:700;background:#f2f2f2}th:nth-child(1){width:30%}th:nth-child(2){width:52%}th:nth-child(3){width:18%}strong{font-weight:700}details{font-size:14px;border-top:1px solid #d8d8d8;margin-top:38px;padding-top:18px;font-family:-apple-system,"PingFang SC",sans-serif}summary{cursor:pointer;font-weight:600}@media(max-width:600px){doc-page{font-size:15px}h1{font-size:23px}th,td{padding:8px 7px;font-size:13px}th:nth-child(1){width:28%}th:nth-child(2){width:52%}th:nth-child(3){width:20%}}@media print{nav,details{display:none}}
@media screen{body{margin:0}}@media screen and (max-width:900px){doc-page{min-width:0;padding:24px 16px;--doc-page-w:min(8.5in,calc(100vw - 32px))!important;--doc-page-margin:24px!important}}
</style></head><body><doc-page margin="0.75in">''' + '\n'.join(content) + notes + '</doc-page><script src="doc-page.js"></script></body></html>\n'
    output = ROOT / "Terms.html"
    output.write_text(page, encoding="utf-8")
    return output


if __name__ == "__main__":
    source_lines, abbreviation_rows = parse_source()
    word, paragraphs = build_word(source_lines, abbreviation_rows)
    preview = build_html(source_lines, abbreviation_rows)
    print(json.dumps({"word": str(word), "html": str(preview), "abbreviations": len(abbreviation_rows) - 1, "key_terms": sum(line.startswith("**") for line in source_lines), "word_paragraphs": paragraphs, "source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(), "word_sha256": hashlib.sha256(word.read_bytes()).hexdigest()}, ensure_ascii=False, indent=2))
